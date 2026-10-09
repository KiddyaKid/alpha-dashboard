#!/usr/bin/env python3
"""Alpha 雷达 v2 — topic heat model + render + publish.

Usage (run from anywhere):
  update_heat.py ingest <topic_id> <counts.json> [--views N] [--post URL]
        counts.json = raw response of x get_posts_counts_recent (granularity=day).
        Writes each UTC day's tweet_count into topic.history; the current (unfinished)
        UTC day is stored with partial_hours so it can be extrapolated.
  update_heat.py set <topic_id> <YYYY-MM-DD> <count> [--partial-hours H] [--views N]
  update_heat.py add-topic <topic.json>     # full topic object (see UPDATE.md)
  update_heat.py retire <topic_id>          # keeps history, hides from ranking
  update_heat.py revive <topic_id>
  update_heat.py event "<title>" "<verdict>" [signal]   # append to 推送历史
  update_heat.py render                     # recompute heat, write index.html + posts.xlsx
  update_heat.py publish ["msg"]            # git add/commit/push (GitHub Pages)
  update_heat.py all ["msg"]                # render + publish
"""
import json, math, sys, subprocess, pathlib
from datetime import datetime, timezone, timedelta
D = pathlib.Path(__file__).resolve().parent
DATA = D / "data.json"
SYD = timezone(timedelta(hours=11))
# Share of a UTC day's X volume that has elapsed after h hours (measured from hourly
# counts on 2026-10-08 UTC); used to extrapolate the unfinished current day.
HOURLY = [89985,90179,93428,95725,85226,83051,80058,73857,75095,84459,93294,110716,
          121561,126176,128521,128419,123990,117881,114068,112997,111951,111298,108247,99587]
def day_fraction(h):
    h = max(0.25, min(24.0, h)); full = int(h); tot = sum(HOURLY)
    return (sum(HOURLY[:full]) + (HOURLY[full] * (h - full) if full < 24 else 0)) / tot

def load(): return json.loads(DATA.read_text())
def save(d): DATA.write_text(json.dumps(d, ensure_ascii=False, indent=1))
def clamp(x, a=0, b=100): return max(a, min(b, x))
def topic(d, tid):
    for t in d["topics"]:
        if t["id"] == tid: return t
    sys.exit(f"no topic {tid}")

def est(h):  # extrapolated full-day count for an unfinished day
    c = h["count"]; ph = h.get("partial_hours")
    return c / day_fraction(ph) if ph else c

def compute(d):
    for t in d["topics"]:
        days = sorted(t["history"])
        ests = [est(t["history"][x]) for x in days]
        # damp early-day noise: blend the extrapolation with yesterday by elapsed share of day
        lh = t["history"][days[-1]]
        if lh.get("partial_hours") and len(days) > 1:
            w = lh["partial_hours"] / 24
            ests[-1] = w * ests[-1] + (1 - w) * ests[-2]
        views = t.get("top_views", 0) or 0
        for i, x in enumerate(days):
            c = max(ests[i], 1)
            vol = clamp((math.log10(c) - 2) / 3 * 100)            # 100/day→0, 100k/day→100
            prev = [e for e in ests[max(0, i-3):i]]
            mom = clamp(50 + 40 * math.log2(c / max(sum(prev)/len(prev), 1))) if prev else 50
            base = 0.625 * vol + 0.375 * mom                       # count-only heat (grid)
            h = t["history"][x]; h["est"] = round(c); h["grid"] = round(base, 1)
        reach = clamp((math.log10(max(views, 1)) - 3) / 4.5 * 100) if views else 0
        last = t["history"][days[-1]]
        t["heat"] = round(0.8 * last["grid"] + 0.2 * reach)
        # trend: mean of last 2 days (today blended) vs mean of the 3 days before them
        recent = ests[-2:]; base = ests[-5:-2] or ests[:1]
        ratio = (sum(recent)/len(recent)) / max(sum(base)/len(base), 1)
        t["delta"] = round((ratio - 1) * 100)                     # % change, 2d vs prior 3d
        t["trend"] = "up" if ratio >= 1.15 else ("down" if ratio <= 0.87 else "flat")
        t["reach"] = round(reach)
    act = sorted([t for t in d["topics"] if t.get("status") == "active"], key=lambda t: -t["heat"])
    for i, t in enumerate(act): t["rank"] = i + 1
    if act and not d["headline"].get("manual"):
        top = act[0]; ups = [t["name"] for t in act if t["trend"] == "up"][:3]
        d["headline"]["text"] = f"热度第一：{top['emoji']} {top['name']}（{top['heat']}）· 升温：" + ("、".join(ups) or "无")
    return d

def render(d):
    compute(d); save(d)
    raw = json.dumps(d, ensure_ascii=False).replace("</", "<\\/")
    (D / "index.html").write_text((D / "template.html").read_text().replace("/*DATA*/", raw))
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        print("openpyxl missing → skip posts.xlsx (use /workspace/.venv-xl/bin/python)"); return
    wb = Workbook(); hdrfill = PatternFill("solid", fgColor="111111")
    def sheet(ws, hdr, rows, link_col=None, widths=()):
        ws.append(hdr)
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = hdrfill
        for r in rows: ws.append(r)
        if link_col:
            for row in ws.iter_rows(min_row=2, min_col=link_col, max_col=link_col):
                for c in row:
                    if c.value: c.hyperlink = c.value; c.font = Font(color="2563EB", underline="single")
        for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws = wb.active; ws.title = "帖子与来源"
    rows, seen = [], set()
    for p in d["posts"]:
        rows.append([p["time"], p["author"], p["text"], p["link"], p["topic"], "█"*p["heat"]+"░"*(10-p["heat"]), p["views"]]); seen.add(p["link"])
    for t in d["topics"]:
        for s in t["sources"]:
            if s["url"] not in seen: rows.append(["", "", s["label"], s["url"], t["name"], "", ""]); seen.add(s["url"])
    sheet(ws, ["时间_悉尼","作者","内容","链接","主题","热度","浏览量"], rows, 4, [12,16,80,52,22,12,12])
    days = sorted({x for t in d["topics"] for x in t["history"]})
    ws2 = wb.create_sheet("话题热度")
    act = sorted(d["topics"], key=lambda t: (t.get("status") != "active", -t["heat"]))
    sheet(ws2, ["排名","话题","状态","热度","趋势","龙头","结论"] + [f"{x[5:]} 帖数" for x in days],
          [[t.get("rank",""), t["name"], t["status"], t["heat"], {"up":"↑","down":"↓"}.get(t["trend"],"→"),
            t["leader"]["sym"], t["verdict"]] + [t["history"].get(x,{}).get("est","") for x in days] for t in act],
          None, [6,28,8,7,6,12,70] + [10]*len(days))
    wb.save(D / "posts.xlsx")
    print("rendered:", ", ".join(f"{t['rank']}.{t['id']}={t['heat']}{'↑' if t['trend']=='up' else '↓' if t['trend']=='down' else '→'}" for t in act if t.get("rank")))

def publish(msg=None):
    msg = msg or "update " + datetime.now(SYD).strftime("%m/%d %H:%M")
    run = lambda *a: subprocess.run(a, cwd=D, check=False)
    run("git", "add", "-A"); r = run("git", "commit", "-m", msg)
    run("git", "push", "origin", "main")

def main(a):
    if not a: print(__doc__); return
    cmd = a[0]; d = load()
    opt = lambda k, default=None: a[a.index(k)+1] if k in a else default
    if cmd == "ingest":
        t = topic(d, a[1]); resp = json.loads(pathlib.Path(a[2]).read_text())
        now = datetime.now(timezone.utc)
        for b in resp["data"]:
            s = datetime.fromisoformat(b["start"].replace("Z","+00:00")); e = datetime.fromisoformat(b["end"].replace("Z","+00:00"))
            hrs = (e - s).total_seconds() / 3600
            if s.hour != 0: continue                      # skip leading partial bucket
            rec = {"count": b["tweet_count"]}
            if hrs < 23.9: rec["partial_hours"] = round(hrs, 2)
            t["history"][s.date().isoformat()] = rec
        day = max(t["history"])
        if opt("--views"): t["top_views"] = int(opt("--views")); t["history"][day]["views"] = t["top_views"]
        if opt("--post"): t["top_post"] = opt("--post")
    elif cmd == "set":
        t = topic(d, a[1]); rec = {"count": int(a[3])}
        if opt("--partial-hours"): rec["partial_hours"] = float(opt("--partial-hours"))
        if opt("--views"): t["top_views"] = int(opt("--views")); rec["views"] = t["top_views"]
        t["history"][a[2]] = rec
    elif cmd == "add-topic":
        nt = json.loads(pathlib.Path(a[1]).read_text()); nt.setdefault("status","active"); nt.setdefault("history",{})
        d["topics"] = [t for t in d["topics"] if t["id"] != nt["id"]] + [nt]
    elif cmd in ("retire", "revive"):
        topic(d, a[1])["status"] = "retired" if cmd == "retire" else "active"
    elif cmd == "event":
        d["history"].append({"time": datetime.now(SYD).strftime("%m/%d %H:%M"), "title": a[1], "verdict": a[2], "signal": a[3] if len(a) > 3 else "info"})
    if cmd in ("ingest","set","add-topic","retire","revive","event"):
        d["meta"]["updated"] = datetime.now(SYD).isoformat(timespec="minutes")
        d["meta"]["updated_label"] = datetime.now(SYD).strftime("%m/%d %H:%M") + " 悉尼"
        save(compute(d)); print("ok"); return
    if cmd == "render": render(d)
    elif cmd == "publish": publish(a[1] if len(a) > 1 else None)
    elif cmd == "all": render(d); publish(a[1] if len(a) > 1 else None)
    else: print(__doc__)

if __name__ == "__main__": main(sys.argv[1:])
