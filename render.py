#!/usr/bin/env python3
"""Rebuild index.html (embeds data.json) and posts.xlsx from data.json."""
import json, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
D = pathlib.Path(__file__).parent
data = json.loads((D/"data.json").read_text())
raw = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
(D/"index.html").write_text((D/"template.html").read_text().replace("/*DATA*/", raw))
wb = Workbook(); ws = wb.active; ws.title = "帖子与来源"
hdr = ["时间_悉尼","作者","内容","链接","主题","热度","浏览量"]; ws.append(hdr)
for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="111111")
links = set()
for p in data["posts"]:
    ws.append([p["time"], p["author"], p["text"], p["link"], p["topic"], "█"*p["heat"]+"░"*(10-p["heat"]), p["views"]]); links.add(p["link"])
for r in data["radar"]:
    for s in r["sources"]:
        if s["url"] not in links: ws.append(["", "", s["label"], s["url"], r["title"], "", ""])
for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
    for c in row: c.hyperlink = c.value; c.font = Font(color="2563EB", underline="single")
for col, w in zip("ABCDEFG", [12,16,80,52,14,12,12]): ws.column_dimensions[col].width = w
wb.save(D/"posts.xlsx")
print("rendered index.html + posts.xlsx;", len(data["posts"]), "posts")
